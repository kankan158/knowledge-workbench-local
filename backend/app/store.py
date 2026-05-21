from __future__ import annotations

import hashlib
import io
import math
import re
from collections import defaultdict
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, Iterable, List, Optional

import chromadb

try:
    from pypdf import PdfReader
except ImportError:
    PdfReader = None

try:
    from docx import Document as DocxDocument
except ImportError:
    DocxDocument = None

TOKEN_PATTERN = re.compile(r"[\w\u4e00-\u9fff]+", re.UNICODE)
CHUNK_TITLE_PATTERN = re.compile(r"\s*\[\d+/\d+\]$")
VECTOR_SIZE = 128


@dataclass
class StoredDocument:
    id: str
    title: str
    content: str
    source: Optional[str] = None


class HashEmbeddingFunction:
    def __call__(self, input: Iterable[str]) -> list[list[float]]:
        return [self._embed_text(text) for text in input]

    def _embed_text(self, text: str) -> list[float]:
        vector = [0.0] * VECTOR_SIZE
        for token in TOKEN_PATTERN.findall(text.lower()):
            digest = hashlib.sha256(token.encode("utf-8")).digest()
            for offset in range(0, 16, 4):
                index = int.from_bytes(digest[offset : offset + 4], "big") % VECTOR_SIZE
                vector[index] += 1.0
        norm = math.sqrt(sum(value * value for value in vector))
        if norm == 0:
            return vector
        return [value / norm for value in vector]


class KnowledgeStore:
    def __init__(self, base_dir: Path) -> None:
        self.base_dir = base_dir
        self.base_dir.mkdir(parents=True, exist_ok=True)
        self.client = chromadb.PersistentClient(path=str(self.base_dir))
        self.embedding_function = HashEmbeddingFunction()
        self.collection = self.client.get_or_create_collection(
            name="knowledge_documents",
            embedding_function=self.embedding_function,
            metadata={"hnsw:space": "cosine"},
        )

    def upsert(self, document: StoredDocument) -> None:
        self.upsert_many([document])

    def upsert_many(self, documents: List[StoredDocument]) -> None:
        if not documents:
            return
        self.collection.upsert(
            ids=[document.id for document in documents],
            documents=[document.content for document in documents],
            metadatas=[
                {
                    "title": document.title,
                    "source": document.source or "",
                    "folder": self._infer_folder_from_source(document.source or ""),
                }
                for document in documents
            ],
        )

    def _infer_folder_from_source(self, source: str) -> str:
        lower_source = source.lower()
        if lower_source.startswith("folder://"):
            # Preserve the full folder path (exclude the filename) so we can match descendants.
            path = source.split("folder://", 1)[1]
            parts = path.split("/")
            if len(parts) <= 1:
                return parts[0].strip() or "unclassified"
            # Remove the trailing filename, keep full folder path like "parent/child"
            folder_path = "/".join(parts[:-1]).strip()
            return folder_path or parts[0].strip() or "unclassified"
        if lower_source.startswith("seed://"):
            return source.split("seed://", 1)[1].split("/", 1)[0].strip() or "seed"
        return "unclassified"

    def _normalize_folder_filter(self, folder_filter: Optional[str]) -> Optional[str]:
        if folder_filter is None:
            return None
        normalized = str(folder_filter).strip().lower()
        return normalized or None

    def _matches_folder_filter(self, metadata: Dict[str, object], folder_filter: Optional[str]) -> bool:
        normalized = self._normalize_folder_filter(folder_filter)
        if not normalized:
            return True

        folder_value = str(metadata.get("folder") or "").lower()
        source_value = str(metadata.get("source") or "").lower()
        title_value = str(metadata.get("title") or "").lower()
        # Match when the normalized filter equals the folder, or the record's folder
        # is a descendant (starts with filter + '/'). Also fall back to substring
        # matching against the source or title for flexibility.
        if normalized == folder_value:
            return True
        if folder_value.startswith(normalized + "/"):
            return True
        if normalized in source_value or normalized in title_value:
            return True
        return False

    def _tokenize(self, text: str) -> List[str]:
        return TOKEN_PATTERN.findall(text.lower())

    def _keyword_score(self, query_tokens: List[str], text: str, title: str = "", source: str = "") -> float:
        if not query_tokens:
            return 0.0

        combined = f"{title}\n{text}\n{source}"
        tokens = self._tokenize(combined)
        if not tokens:
            return 0.0

        token_counts = defaultdict(int)
        for token in tokens:
            token_counts[token] += 1

        overlap = 0.0
        for token in query_tokens:
            if token in token_counts:
                overlap += 1.0 + min(token_counts[token], 3) * 0.15

        title_bonus = 0.0
        normalized_title = title.lower()
        for token in query_tokens:
            if token in normalized_title:
                title_bonus += 0.25

        score = (overlap / max(len(query_tokens), 1)) * 0.75 + title_bonus
        return min(score, 1.0)

    def _base_title(self, title: str) -> str:
        return CHUNK_TITLE_PATTERN.sub("", title).strip() or title

    def _collect_records(self) -> List[Dict[str, object]]:
        result = self.collection.get(include=["documents", "metadatas"])
        records: list[dict[str, object]] = []
        ids = result.get("ids", [])
        documents = result.get("documents", [])
        metadatas = result.get("metadatas", [])
        for index, document_id in enumerate(ids):
            metadata = metadatas[index] if index < len(metadatas) else {}
            content = documents[index] if index < len(documents) else ""
            records.append(
                {
                    "id": document_id,
                    "title": metadata.get("title", "Untitled"),
                    "content": content,
                    "source": metadata.get("source") or None,
                    "folder": metadata.get("folder") or None,
                }
            )
        return records

    def _score_records(
        self,
        query: str,
        search_type: str = "hybrid",
        alpha: float = 0.5,
        folder_filter: Optional[str] = None,
    ) -> List[Dict[str, object]]:
        records = self._collect_records()
        normalized_query = query.strip()
        query_tokens = self._tokenize(normalized_query)
        if not normalized_query:
            return []

        def build_scores(active_folder_filter: Optional[str]) -> List[Dict[str, object]]:
            vector_scores: dict[str, float] = {}
            keyword_scores: dict[str, float] = {}

            if search_type in {"vector", "hybrid"}:
                vector_result = self.collection.query(
                    query_texts=[query],
                    n_results=min(max(len(records), 1), 50),
                    include=["documents", "metadatas", "distances"],
                )
                distances = vector_result.get("distances", [[]])[0]
                ids = vector_result.get("ids", [[]])[0]
                for index, document_id in enumerate(ids):
                    distance = distances[index] if index < len(distances) else None
                    if distance is None:
                        continue
                    vector_scores[str(document_id)] = max(0.0, 1.0 - float(distance))

            if search_type in {"keyword", "hybrid"}:
                for record in records:
                    keyword_scores[str(record["id"])] = self._keyword_score(
                        query_tokens,
                        str(record.get("content") or ""),
                        title=str(record.get("title") or ""),
                        source=str(record.get("source") or ""),
                    )

            combined: list[dict[str, object]] = []
            for record in records:
                if not self._matches_folder_filter(
                    {"folder": record.get("folder"), "source": record.get("source"), "title": record.get("title")},
                    active_folder_filter,
                ):
                    continue

                vector_score = vector_scores.get(str(record["id"]), 0.0)
                keyword_score = keyword_scores.get(str(record["id"]), 0.0)

                if search_type == "vector":
                    score = vector_score
                elif search_type == "keyword":
                    score = keyword_score
                else:
                    score = (alpha * vector_score) + ((1.0 - alpha) * keyword_score)

                combined.append({**record, "score": score})

            combined.sort(key=lambda item: float(item.get("score") or 0.0), reverse=True)
            return combined

        combined = build_scores(folder_filter)
        if folder_filter and not combined:
            combined = build_scores(None)

        grouped: dict[str, dict[str, object]] = {}
        for record in combined:
            source = str(record.get("source") or "")
            title = self._base_title(str(record.get("title") or "Untitled"))
            group_key = f"{source}||{title}"

            current = grouped.get(group_key)
            if not current:
                grouped[group_key] = {
                    **record,
                    "title": title,
                    "content": str(record.get("content") or ""),
                    "score": float(record.get("score") or 0.0),
                }
                continue

            current_score = float(current.get("score") or 0.0)
            next_score = float(record.get("score") or 0.0)
            if next_score > current_score:
                current["id"] = record.get("id")
                current["content"] = record.get("content")
                current["score"] = next_score

        deduped = list(grouped.values())
        deduped.sort(key=lambda item: float(item.get("score") or 0.0), reverse=True)
        return deduped

    def chunk_text(self, text: str, max_chars: int = 1200, overlap: int = 150) -> List[str]:
        normalized = text.replace("\r\n", "\n").strip()
        if not normalized:
            return []

        paragraphs = [paragraph.strip() for paragraph in normalized.split("\n\n") if paragraph.strip()]
        if not paragraphs:
            paragraphs = [normalized]

        chunks: list[str] = []
        current = ""
        for paragraph in paragraphs:
            candidate = f"{current}\n\n{paragraph}" if current else paragraph
            if len(candidate) <= max_chars:
                current = candidate
                continue

            if current:
                chunks.append(current)
                current = ""

            if len(paragraph) <= max_chars:
                current = paragraph
                continue

            start = 0
            while start < len(paragraph):
                end = min(len(paragraph), start + max_chars)
                if end < len(paragraph):
                    split_at = paragraph.rfind(" ", start, end)
                    if split_at > start + max_chars // 2:
                        end = split_at
                piece = paragraph[start:end].strip()
                if piece:
                    chunks.append(piece)
                if end >= len(paragraph):
                    break
                start = max(end - overlap, start + 1)

        if current:
            chunks.append(current)

        return chunks

    def ingest_text(self, base_id: str, title: str, content: str, source: Optional[str] = None) -> List[StoredDocument]:
        chunks = self.chunk_text(content)
        if not chunks:
            chunks = [content.strip()]

        documents: list[StoredDocument] = []
        chunk_total = len(chunks)
        for index, chunk in enumerate(chunks):
            digest = hashlib.sha1(chunk.encode("utf-8")).hexdigest()[:12]
            chunk_title = title if chunk_total == 1 else f"{title} [{index + 1}/{chunk_total}]"
            documents.append(
                StoredDocument(
                    id=f"{base_id}:{index:04d}:{digest}",
                    title=chunk_title,
                    content=chunk,
                    source=source,
                )
            )

        self.upsert_many(documents)
        return documents

    def extract_pdf_text(self, file_bytes: bytes) -> str:
        if not PdfReader:
            raise ImportError("pypdf is not installed. Install it with: pip install pypdf")
        try:
            pdf_reader = PdfReader(io.BytesIO(file_bytes))
            text_parts: list[str] = []
            for page in pdf_reader.pages:
                extracted = page.extract_text()
                if extracted:
                    text_parts.append(extracted)
            return "\n\n".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to extract PDF: {e}")

    def extract_docx_text(self, file_bytes: bytes) -> str:
        if not DocxDocument:
            raise ImportError("python-docx is not installed. Install it with: pip install python-docx")
        try:
            doc = DocxDocument(io.BytesIO(file_bytes))
            text_parts: list[str] = []
            for paragraph in doc.paragraphs:
                if paragraph.text.strip():
                    text_parts.append(paragraph.text)
            return "\n\n".join(text_parts)
        except Exception as e:
            raise ValueError(f"Failed to extract DOCX: {e}")

    def extract_text_from_file(
        self, file_bytes: bytes, filename: str
    ) -> str:
        lower_name = filename.lower()
        if lower_name.endswith(".pdf"):
            return self.extract_pdf_text(file_bytes)
        elif lower_name.endswith((".docx", ".doc")):
            return self.extract_docx_text(file_bytes)
        else:
            try:
                return file_bytes.decode("utf-8-sig")
            except UnicodeDecodeError:
                return file_bytes.decode("gb18030", errors="ignore")

    def search(
        self,
        query: str,
        top_k: int = 5,
        min_score: float = 0.05,
        preview_chars: int = 400,
        include_full: bool = False,
        search_type: str = "hybrid",
        alpha: float = 0.5,
        folder_filter: Optional[str] = None,
    ) -> List[Dict[str, object]]:
        """Perform vector, keyword, or hybrid search and return results.

        - `min_score` filters out low-similarity results (0.0-1.0).
        - `preview_chars` limits the returned `content` length unless `include_full` is True.
        - `include_full` when True returns the full chunk content (used by /ask to synthesize).
        """
        mode = (search_type or "hybrid").strip().lower()
        if mode not in {"vector", "keyword", "hybrid"}:
            mode = "hybrid"

        scored_records = self._score_records(
            query=query,
            search_type=mode,
            alpha=max(0.0, min(float(alpha), 1.0)),
            folder_filter=folder_filter,
        )

        items: list[dict[str, object]] = []
        for record in scored_records:
            score = float(record.get("score") or 0.0)
            if min_score is not None and score < float(min_score):
                continue

            content = str(record.get("content") or "")
            if include_full:
                returned_content = content
            else:
                returned_content = content[:preview_chars]
                if len(content) > preview_chars:
                    returned_content = returned_content.rstrip()
                    returned_content += "..."

            items.append(
                {
                    "id": record.get("id"),
                    "title": record.get("title", "Untitled"),
                    "content": returned_content,
                    "source": record.get("source") or None,
                    "score": score,
                }
            )
            if len(items) >= top_k:
                break

        if folder_filter and not items:
            return self.search(
                query,
                top_k=top_k,
                min_score=min_score,
                preview_chars=preview_chars,
                include_full=include_full,
                search_type=mode,
                alpha=alpha,
                folder_filter=None,
            )
        return items

    def list_all(self) -> List[Dict[str, object]]:
        result = self.collection.get(include=["documents", "metadatas"])
        items: list[dict[str, object]] = []
        ids = result.get("ids", [])
        documents = result.get("documents", [])
        metadatas = result.get("metadatas", [])
        for index, document_id in enumerate(ids):
            metadata = metadatas[index] if index < len(metadatas) else {}
            content = documents[index] if index < len(documents) else ""
            items.append(
                {
                    "id": document_id,
                    "title": metadata.get("title", "Untitled"),
                    "content": content,
                    "source": metadata.get("source") or None,
                    "score": None,
                }
            )
        return items

    def count(self) -> int:
        return self.collection.count()

    def delete_document(self, document_id: str) -> None:
        self.collection.delete(ids=[document_id])

    def seed(self, documents: List[StoredDocument]) -> None:
        existing_ids = set(self.collection.get(include=[])["ids"])
        for document in documents:
            if document.id not in existing_ids:
                self.upsert(document)
